"""Extracción de texto desde bytes en memoria, con límites defensivos. No escribe a disco."""
from __future__ import annotations

import io
import zipfile
from pathlib import PurePath

from defusedxml import ElementTree as DefusedET

from features.preanalisis.domain.models import (
    MAX_BYTES_ARCHIVO,
    MAX_BYTES_DOCUMENT_XML,
    MAX_CARACTERES_TEXTO,
    MAX_PAGINAS_PDF,
    ArchivoDemasiadoGrande,
    DocumentoIlegible,
    ExtraccionResultado,
    FormatoNoSoportado,
    Segmento,
)

_EXT_TEXTO = {".txt", ".md", ".csv"}
_EXT_IMAGEN = {".png", ".jpg", ".jpeg"}
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_PARRAFOS_POR_BLOQUE = 10
_MAX_FILAS_HOJA = 50_000
MAX_PARRAFOS_DOCX = 100_000
MAX_CELDAS_XLSX = 500_000
MAX_HOJAS_XLSX = 50
MAX_MIEMBRO_ZIP = 30 * 1024 * 1024
MAX_TOTAL_ZIP = 120 * 1024 * 1024
MAX_MIEMBROS_ZIP = 2000
RAZON_COMPRESION_MAX = 100
_UMBRAL_RAZON = 1024 * 1024  # la razón solo se evalúa en miembros grandes; los pequeños comprimen "de más" sin riesgo


def verificar_zip(contenido: bytes) -> None:
    """Rechaza zip bombs por cabeceras (tamaño por miembro, total, razón de compresión) sin descomprimir nada."""
    with zipfile.ZipFile(io.BytesIO(contenido)) as z:
        miembros = z.infolist()
    if len(miembros) > MAX_MIEMBROS_ZIP:
        raise DocumentoIlegible("El contenedor del documento tiene demasiados elementos.")
    total = 0
    for m in miembros:
        total += m.file_size
        if m.file_size > MAX_MIEMBRO_ZIP or total > MAX_TOTAL_ZIP:
            raise DocumentoIlegible("El contenido descomprimido del documento excede el límite.")
        if m.file_size > _UMBRAL_RAZON and m.file_size > RAZON_COMPRESION_MAX * max(m.compress_size, 1):
            raise DocumentoIlegible("El documento tiene una razón de compresión sospechosa.")


def _extension(nombre: str) -> str:
    return PurePath(nombre or "").suffix.lower()


def extraer_texto(nombre: str, contenido: bytes) -> ExtraccionResultado:
    ext = _extension(nombre)
    if ext in _EXT_IMAGEN:
        raise FormatoNoSoportado("Las imágenes no son analizables (no hay OCR).")
    if ext not in ({".pdf", ".docx", ".xlsx"} | _EXT_TEXTO):
        raise FormatoNoSoportado(f"Formato no soportado: '{ext or 'sin extensión'}'.")
    if len(contenido) > MAX_BYTES_ARCHIVO:
        raise ArchivoDemasiadoGrande("El archivo excede 25 MB.")
    if not contenido:
        raise DocumentoIlegible("El archivo está vacío.")
    try:
        if ext == ".pdf":
            segmentos = _pdf(contenido)
        elif ext == ".docx":
            segmentos = _docx(contenido)
        elif ext == ".xlsx":
            segmentos = _xlsx(contenido)
        else:
            segmentos = _texto(contenido)
    except (FormatoNoSoportado, ArchivoDemasiadoGrande, DocumentoIlegible):
        raise
    except Exception as exc:  # noqa: BLE001 - cualquier fallo del parser = documento ilegible
        raise DocumentoIlegible("No se pudo leer el documento (dañado, cifrado o inválido).") from exc

    total = 0
    acotados: list[Segmento] = []
    truncado = False
    for seg in segmentos:
        if not seg.texto.strip():
            continue
        restante = MAX_CARACTERES_TEXTO - total
        if restante <= 0:
            truncado = True
            break
        texto = seg.texto[:restante]
        truncado = truncado or len(texto) < len(seg.texto)
        acotados.append(Segmento(seg.origen, texto))
        total += len(texto)
    if not acotados:
        raise DocumentoIlegible("El documento no contiene texto extraíble.")
    return ExtraccionResultado(ext.lstrip("."), acotados, total, truncado)


def _pdf(contenido: bytes) -> list[Segmento]:
    from pypdf import PdfReader

    lector = PdfReader(io.BytesIO(contenido))
    if lector.is_encrypted:
        raise DocumentoIlegible("El PDF está cifrado.")
    if len(lector.pages) > MAX_PAGINAS_PDF:
        raise DocumentoIlegible(f"El PDF excede {MAX_PAGINAS_PDF} páginas.")
    salida: list[Segmento] = []
    total = 0
    for n, pagina in enumerate(lector.pages, start=1):
        texto = pagina.extract_text() or ""
        salida.append(Segmento(f"p. {n}", texto))
        total += len(texto)
        if total > MAX_CARACTERES_TEXTO:
            break
    return salida


def _docx(contenido: bytes) -> list[Segmento]:
    verificar_zip(contenido)
    with zipfile.ZipFile(io.BytesIO(contenido)) as z:
        try:
            info = z.getinfo("word/document.xml")
        except KeyError as exc:
            raise DocumentoIlegible("El .docx no contiene word/document.xml.") from exc
        if info.file_size > MAX_BYTES_DOCUMENT_XML:
            raise DocumentoIlegible("El contenido descomprimido del .docx excede el límite.")
        with z.open(info) as f:
            xml = f.read(MAX_BYTES_DOCUMENT_XML + 1)
    if len(xml) > MAX_BYTES_DOCUMENT_XML:
        raise DocumentoIlegible("El contenido descomprimido del .docx excede el límite.")
    parrafos: list[str] = []
    partes: list[str] = []
    n_parrafos = 0
    caracteres = 0
    tag_p, tag_t = f"{_W}p", f"{_W}t"
    tags_espacio = (f"{_W}tab", f"{_W}br")
    # Streaming: se aborta al superar los topes sin construir el árbol completo en memoria.
    for _ev, el in DefusedET.iterparse(io.BytesIO(xml), events=("end",)):
        if el.tag == tag_t:
            if el.text:
                partes.append(el.text)
        elif el.tag in tags_espacio:
            partes.append(" ")
        elif el.tag == tag_p:
            n_parrafos += 1
            if n_parrafos > MAX_PARRAFOS_DOCX:
                raise DocumentoIlegible("El .docx excede el número máximo de párrafos.")
            linea = "".join(partes).strip()
            partes = []
            if linea:
                parrafos.append(linea)
                caracteres += len(linea)
                if caracteres > MAX_CARACTERES_TEXTO:
                    break
            el.clear()
    salida: list[Segmento] = []
    for i in range(0, len(parrafos), _PARRAFOS_POR_BLOQUE):
        n = i // _PARRAFOS_POR_BLOQUE + 1
        salida.append(Segmento(f"bloque {n}", "\n".join(parrafos[i : i + _PARRAFOS_POR_BLOQUE])))
    return salida


def _xlsx(contenido: bytes) -> list[Segmento]:
    from openpyxl import load_workbook

    verificar_zip(contenido)  # incluye xl/sharedStrings.xml, que openpyxl carga completo
    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    salida: list[Segmento] = []
    total = 0
    celdas_vistas = 0
    try:
        if len(libro.worksheets) > MAX_HOJAS_XLSX:
            raise DocumentoIlegible("El libro tiene demasiadas hojas.")
        for hoja in libro.worksheets:
            if total > MAX_CARACTERES_TEXTO or celdas_vistas > MAX_CELDAS_XLSX:
                break
            filas: list[str] = []
            for n, fila in enumerate(hoja.iter_rows(values_only=True)):
                celdas_vistas += len(fila)
                if n >= _MAX_FILAS_HOJA or total > MAX_CARACTERES_TEXTO or celdas_vistas > MAX_CELDAS_XLSX:
                    break
                celdas = [str(c).strip() for c in fila if c is not None and str(c).strip()]
                if celdas:
                    linea = " | ".join(celdas)
                    filas.append(linea)
                    total += len(linea)
            salida.append(Segmento(f"hoja {hoja.title}", "\n".join(filas)))
    finally:
        libro.close()
    return salida


def _texto(contenido: bytes) -> list[Segmento]:
    try:
        texto = contenido.decode("utf-8-sig")
    except UnicodeDecodeError:
        texto = contenido.decode("latin-1")
    return [Segmento("texto", texto)]
